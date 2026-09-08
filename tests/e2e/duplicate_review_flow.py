"""Real queued scan → reviewed workbook snapshot → reusable template."""
import json
from contextlib import closing
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook
from playwright.sync_api import expect

from api.core.file_duplicates import _stable_hash


def verify_duplicate_review(api, page, root, report_dir):
    source = root / 'duplicate-fixtures'
    output = root / 'duplicate-reports'
    files = [('alpha/report.txt', 'same'), ('beta/report.txt', 'same'),
             ('gamma/report.txt', 'different-content'), ('unique.txt', 'single')]
    for name, text in files:
        path = source / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding='utf-8')
    navigation = page.get_by_role('navigation', name='工具', exact=True)
    navigation.get_by_role('button', name='文件批处理', exact=True).click()
    page.get_by_role('navigation', name='子功能', exact=True).get_by_role('button', name='重复审查', exact=True).click()
    panel = page.locator('.dedup-panel')
    page.evaluate('path => { window.__duplicatePicker = window.pywebview.api.system_pySelectDirDialog; window.pywebview.api.system_pySelectDirDialog = async () => path }', str(source))
    try:
        panel.get_by_role('button', name='选择目录', exact=True).click()
        start = panel.get_by_role('button', name='开始扫描', exact=True)
        before = {task['id'] for task in api.task_list()['tasks']}
        start.click()
        expect(panel.locator('.dedup-summary')).to_contain_text('扫描完整 · 1 组相同内容')
        expect(panel.locator('.duplicate-files li')).to_have_count(2)
        new_tasks = [task for task in api.task_list()['tasks'] if task['id'] not in before]
        assert len(new_tasks) == 1 and new_tasks[0]['method'] == 'file_deduplicate', new_tasks
        assert new_tasks[0]['result']['duplicateBytes'] == 4
        assert new_tasks[0]['result']['hashedFiles'] == 2
        panel.locator('.el-radio-button').filter(has_text='同名候选').click()
        expect(panel.locator('.dedup-summary')).to_have_count(0)
        start.click()
        expect(panel.locator('.dedup-summary')).to_contain_text('扫描完整 · 1 组同名候选')
        expect(panel.locator('.dedup-summary')).not_to_contain_text('逻辑体积')
        expect(panel.locator('.duplicate-files li')).to_have_count(3)
        panel.locator('.el-radio-button').filter(has_text='比较内容').click()
        start.click()
        expect(panel.locator('.dedup-summary')).to_contain_text('1 组相同内容')
        page.evaluate('path => { window.pywebview.api.system_pySelectDirDialog = async () => path }', str(output))
        panel.get_by_role('button', name='选择报告目录', exact=True).click()
        export = panel.get_by_role('button', name='导出 Excel 报告', exact=True)
        export.click()
        handoff = panel.get_by_role('button', name='检查结果 / 继续处理（1）', exact=True)
        expect(handoff).to_be_visible()
        reports = list(output.glob('*.xlsx'))
        assert len(reports) == 1
        with closing(load_workbook(reports[0], read_only=True)) as book:
            rows = list(book['逐文件核对'].values)
            assert len(rows) == 3
            assert all(row[3] != '未比较内容' for row in rows[1:])
            assert book['扫描说明']['B2'].value == '完整'
        assert all((source / name).read_text(encoding='utf-8') == text for name, text in files)
        page.set_viewport_size({'width': 980, 'height': 760})
        panel.locator('.dedup-summary').evaluate("node => node.scrollIntoView({block:'start'})")
        expect(page.locator('.el-message')).to_have_count(0)
        for theme in ('light', 'dark'):
            page.evaluate("theme => { document.documentElement.dataset.theme = theme; document.documentElement.classList.toggle('dark', theme === 'dark') }", theme)
            page.wait_for_timeout(200)
            assert panel.evaluate('(node) => node.scrollWidth <= node.clientWidth')
            page.screenshot(path=str(report_dir / f'duplicate-review-{theme}.png'))
        page.set_viewport_size({'width': 1440, 'height': 1000})
        page.evaluate("document.documentElement.dataset.theme='light'; document.documentElement.classList.remove('dark')")
        before = {task['id'] for task in api.task_list()['tasks']}
        handoff.click()
        dialog = page.get_by_role('dialog', name='处理结果', exact=True)
        dialog.locator('.route-select').click()
        expect(page.get_by_role('option').filter(has_text='数据质检')).to_be_visible()
        assert {task['id'] for task in api.task_list()['tasks']} == before
        page.keyboard.press('Escape')
        dialog.get_by_role('button', name='取消', exact=True).click()

        # Actual template resolves structured groups and summary without saving
        # a second copy of every file path inside the summary.
        navigation.get_by_role('button', name='自动化工作流', exact=True).click()
        page.get_by_role('tab', name='工作流', exact=True).click()
        page.locator('.template-card').filter(has_text='重复文件扫描 → 核对报告').get_by_role('button', name='使用模板', exact=True).click()
        editor = page.locator('.workflow-editor')
        inputs = editor.locator('.run-setup')
        inputs.locator('[data-field-name="directory"] input').fill(str(source))
        inputs.locator('[data-field-name="outputDir"] input').fill(str(output))
        run = editor.get_by_role('button', name='立即运行', exact=True)
        run.click()
        workflow = next(item for item in api.workflow_list()['workflows'] if item['name'].startswith('重复文件扫描'))
        completed = None
        for _ in range(400):
            completed = next((item for item in api.workflow_list()['runs'] if item['workflowId'] == workflow['id']), None)
            if completed and completed['status'] != 'running':
                break
            page.wait_for_timeout(100)
        assert completed and completed['status'] == 'success', completed
        assert Path(completed['steps'][1]['result']['file']).is_file()
        expect(page.get_by_role('tab', name='运行记录', exact=True)).to_have_attribute('aria-selected', 'true')

        # Verify the actual persisted resumed record shown to the user, including
        # replacing a downstream report after a partial query changed its matches.
        retry_source = root / 'duplicate-resume'
        retry_source.mkdir()
        for name in ('a.txt', 'b.txt', 'c.txt'):
            (retry_source / name).write_bytes(b'abc')
        retry_workflow = api.workflow_create_from_template({'templateId': 'builtin-duplicate-report'})['workflow']
        retry_workflow['steps'][0]['onPartial'] = 'continue'
        retry_workflow = api.workflow_save(retry_workflow)['workflow']
        retry_inputs = {'directory': str(retry_source), 'outputDir': str(output), 'extensions': [], 'recursive': True, 'limit': 100}
        def temporary_failure(path, expected, check):
            if path.name == 'c.txt':
                raise PermissionError('temporary fixture failure')
            return _stable_hash(path, expected, check)
        with mock.patch('api.core.file_duplicates._stable_hash', side_effect=temporary_failure):
            partial_run = api.workflow_run({'id': retry_workflow['id'], 'input': retry_inputs})
        assert partial_run['run']['status'] == 'partial', partial_run
        old_report = Path(partial_run['context']['steps']['report']['file'])
        (retry_source / 'b.txt').write_bytes(b'now different')
        (retry_source / 'c.txt').write_bytes(b'also unique contents')
        resumed = api.workflow_run({'id': retry_workflow['id'], 'input': retry_inputs, '_resumeRunId': partial_run['run']['id']})
        assert resumed['run']['status'] == 'success', resumed
        assert resumed['context']['steps']['scan']['outputPaths'] == []
        assert Path(resumed['context']['steps']['report']['file']) != old_report
        page.get_by_role('button', name='刷新', exact=True).click()
        matching_records = page.locator('.run-list .el-collapse-item').filter(has_text=retry_workflow['name'])
        expect(matching_records).to_have_count(3)
        record = matching_records.first
        record.locator('.el-collapse-item__header').click()
        expect(record.get_by_text('前序依赖已重新执行，本步骤使用当前结果重新生成。', exact=False)).to_be_visible()
        page.set_viewport_size({'width': 980, 'height': 760})
        record.locator('.resume-info').last.scroll_into_view_if_needed()
        expect(page.locator('.el-message')).to_have_count(0)
        page.screenshot(path=str(report_dir / 'duplicate-report-resumed.png'))
        page.set_viewport_size({'width': 1440, 'height': 1000})

        # The explicit scan cap is reflected in the UI and in an exported
        # partial snapshot; a large group never mounts more than 20 file rows.
        for index in range(101):
            (source / f'many-{index}.txt').write_text('same', encoding='utf-8')
        navigation.get_by_role('button', name='文件批处理', exact=True).click()
        page.get_by_role('navigation', name='子功能', exact=True).get_by_role('button', name='重复审查', exact=True).click()
        limit = panel.locator('.el-form-item').filter(has_text='最多扫描文件').locator('input')
        limit.fill('100')
        limit.press('Tab')
        start.click()
        expect(panel.locator('.dedup-summary')).to_contain_text('扫描不完整')
        expect(panel.locator('.el-alert')).to_contain_text('达到文件数量上限')
        expect(panel.locator('.duplicate-files li')).to_have_count(20)
        panel.get_by_label('组内文件分页').locator('.btn-next').click()
        expect(panel.locator('.duplicate-files li')).to_have_count(20)
        export.click()
        expect(handoff).to_be_visible()
        latest = max(output.glob('*.xlsx'), key=lambda path: path.stat().st_mtime_ns)
        with closing(load_workbook(latest, read_only=True)) as book:
            assert '不完整' in book['扫描说明']['B2'].value
        panel.locator('.dedup-summary').evaluate("node => node.scrollIntoView({block:'start'})")
        page.set_viewport_size({'width': 980, 'height': 760})
        expect(page.locator('.el-message')).to_have_count(0)
        page.screenshot(path=str(report_dir / 'duplicate-review-partial.png'))
        extensions = panel.locator('.el-form-item').filter(has_text='扩展名').locator('input')
        extensions.fill('no-match')
        expect(panel.locator('.dedup-summary')).to_have_count(0)
        expect(handoff).to_have_count(0)
        start.click()
        expect(panel.get_by_text('当前范围内未发现重复候选', exact=True)).to_be_visible()
    finally:
        page.evaluate('() => { window.pywebview.api.system_pySelectDirDialog = window.__duplicatePicker; delete window.__duplicatePicker }')
        page.set_viewport_size({'width': 1440, 'height': 1000})
        page.evaluate("document.documentElement.dataset.theme='light'; document.documentElement.classList.remove('dark')")
    result = {'passed': True, 'checks': [
        'real scan enters persistent task queue and skips unique-size hashing',
        'name mode never claims duplicate bytes; mode changes clear old results',
        'real workbook snapshot preserves sources and connects to Excel tools',
        'template binds groups and summary and creates actual workbook',
        'partial query retry clears old matches and regenerates the dependent report with visible reason',
        '100-file cap is explicit and partial workbook is labelled',
        'large group renders only 20 files per page; empty and stale states',
        '980x760 light/dark and partial review evidence']}
    (report_dir / 'duplicate-review.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return result
