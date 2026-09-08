import os
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from api.core.context import TaskCancelled, TaskContext, task_context
from api.core.file_duplicates import _stable_hash, export_duplicate_report, scan_duplicates
from api.file import FileTool


class FileDuplicateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / 'source'
        self.source.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def make(self, name, content=b'abc'):
        path = self.source / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        return path

    def scan(self, **kwargs):
        return scan_duplicates(self.source, {'recursive': True}, 100, **kwargs)

    def test_only_same_size_candidates_are_hashed_and_content_is_verified(self):
        first = self.make('a.txt')
        second = self.make('nested/b.txt')
        self.make('different.txt', b'xyz')
        self.make('unique.txt', b'longer')
        with mock.patch('api.core.file_duplicates._stable_hash', wraps=_stable_hash) as hashing:
            result = self.scan()
        self.assertEqual(hashing.call_count, 3)
        self.assertTrue(result['complete'])
        self.assertEqual(result['totalGroups'], 1)
        self.assertEqual(set(result['groups'][0]['files']), {str(first.resolve()), str(second.resolve())})
        self.assertEqual(result['duplicateBytes'], 3)
        self.assertEqual(result['hashedFiles'], 3)
        self.assertEqual(result['candidateFiles'], 3)
        self.assertTrue(result['summary']['scanId'])

    def test_name_candidates_do_not_claim_content_or_reclaimable_space(self):
        self.make('a/same.txt', b'short')
        self.make('b/same.txt', b'completely different')
        with mock.patch('api.core.file_duplicates._stable_hash') as hashing:
            result = self.scan(mode='name')
        hashing.assert_not_called()
        self.assertEqual(result['totalGroups'], 1)
        self.assertIsNone(result['duplicateBytes'])
        self.assertIsNone(result['spaceSaved'])
        self.assertIsNone(result['groups'][0]['digest'])

    def test_limit_exact_and_extra_match_via_api(self):
        for index in range(100):
            self.make(f'{index}.txt')
        result = FileTool().file_deduplicate({'directory': str(self.source), 'limit': 100})
        self.assertTrue(result['complete'])
        self.make('extra.txt')
        result = FileTool().file_deduplicate({'directory': str(self.source), 'limit': 100})
        self.assertTrue(result['truncated'])
        self.assertEqual(result['scanned'], 100)
        self.assertFalse(result['complete'])
        self.assertEqual(result['limit'], 100)

    def test_hash_failure_keeps_confirmed_groups_but_marks_partial(self):
        self.make('a.txt')
        self.make('b.txt')
        failed = self.make('c.txt')
        def read(path, expected, check):
            if path == failed:
                raise PermissionError('no read access')
            return _stable_hash(path, expected, check)
        with mock.patch('api.core.file_duplicates._stable_hash', side_effect=read):
            result = self.scan()
        self.assertTrue(result['partial'])
        self.assertEqual(result['totalGroups'], 1)
        self.assertEqual(result['errorCount'], 1)
        self.assertEqual(result['hashedFiles'], 2)

    def test_hash_budget_and_cancellation_are_not_silent_success(self):
        self.make('a.txt')
        self.make('b.txt')
        result = self.scan(max_hash_seconds=0)
        self.assertTrue(result['partial'])
        self.assertEqual(result['errors'][0]['code'], 'DUPLICATE_HASH_TIME_LIMIT')
        context = TaskContext()
        context.cancel.set()
        with task_context(context), self.assertRaises(TaskCancelled):
            self.scan()

    def test_change_while_hashing_is_rejected(self):
        path = self.make('a.txt')
        calls = 0
        def change():
            nonlocal calls
            calls += 1
            if calls == 1:
                path.write_bytes(b'longer changed contents')
        with self.assertRaisesRegex(ValueError, '变化'):
            _stable_hash(path, 3, change)

    def test_cancellation_during_comparison_reaches_api_caller(self):
        self.make('a.txt')
        self.make('b.txt')
        context = TaskContext()
        context.callback = lambda event: context.cancel.set() if event.get('current', 0) > 0 else None
        with task_context(context), self.assertRaises(TaskCancelled):
            FileTool().file_deduplicate({'directory': str(self.source)})
        self.assertEqual(context.outputs, [])

    def test_links_are_excluded_and_hardlinks_do_not_inflate_bytes(self):
        first = self.make('a.txt')
        try:
            os.link(first, self.source / 'alias.txt')
        except OSError:
            self.skipTest('filesystem does not support hardlinks')
        self.make('copy.txt')
        result = self.scan()
        self.assertEqual(result['groups'][0]['independentCopies'], 2)
        self.assertEqual(result['hardLinkAliases'], 1)
        self.assertEqual(result['duplicateBytes'], 3)

    def test_scope_recursive_extensions_and_excluded_output(self):
        self.make('a.txt')
        self.make('nested/b.txt')
        self.make('delivery/old.txt')
        self.make('other.csv')
        self.make('.ppx_history/old.txt')
        result = scan_duplicates(self.source, {'recursive': True, 'extensions': ['txt'], 'exclude_directory': self.source / 'delivery'}, 100)
        self.assertEqual(result['scanned'], 2)
        result = scan_duplicates(self.source, {'recursive': False, 'extensions': ['txt']}, 100)
        self.assertEqual(result['scanned'], 1)
        self.assertEqual(result['totalGroups'], 0)

    def test_invalid_mode_and_exclusion_are_rejected(self):
        for opts in ({'mode': 'garbage'}, {'excludeDirectory': str(self.source)}):
            result = FileTool().file_deduplicate({'directory': str(self.source), **opts})
            self.assertNotEqual(result['code'], 0)

    def test_report_is_snapshot_formula_safe_and_does_not_overwrite(self):
        first, second = self.make('a.txt'), self.make('b.txt')
        result = self.scan()
        # Formula-looking user values must remain text in the exported workbook.
        result['summary']['scope']['keyword'] = '=1+1'
        result['groups'][0]['files'][0] = '=HYPERLINK("https://example.invalid")'
        first.unlink()
        target = export_duplicate_report(result['groups'], result['summary'], self.root / 'reports')
        second_target = export_duplicate_report(result['groups'], result['summary'], self.root / 'reports')
        self.assertNotEqual(target, second_target)
        with closing(load_workbook(target, read_only=True)) as book:
            self.assertEqual(book['逐文件核对']['B2'].data_type, 's')
            self.assertTrue(book['逐文件核对']['B2'].value.startswith('=HYPERLINK'))
            self.assertIn('扫描时快照', book['扫描说明']['B1'].value)
            self.assertEqual(book['扫描说明']['B2'].value, '完整')
        self.assertEqual(second.read_bytes(), b'abc')

    def test_partial_and_empty_snapshot_reports_are_explicit(self):
        result = self.scan()
        result['summary'].update(complete=False, partial=True)
        target = export_duplicate_report([], result['summary'], self.root / 'reports')
        with closing(load_workbook(target, read_only=True)) as book:
            self.assertIn('不完整', book['扫描说明']['B2'].value)
            self.assertEqual(len(list(book['逐文件核对'].rows)), 1)

    def test_invalid_report_blocks_file_creation(self):
        result = self.scan()
        result['summary']['totalGroups'] = 2
        with self.assertRaisesRegex(ValueError, '数量'):
            export_duplicate_report([], result['summary'], self.root / 'reports')
        self.assertFalse((self.root / 'reports').exists())
        with self.assertRaises(ValueError):
            export_duplicate_report([], {}, self.root / 'reports')


if __name__ == '__main__':
    unittest.main()
