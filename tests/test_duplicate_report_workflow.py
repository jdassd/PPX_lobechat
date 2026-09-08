import tempfile
import unittest
import zipfile
from contextlib import closing
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from api.api import API
from api.core.file_duplicates import _stable_hash
from api.workflow import WorkflowMixin, _depends_on_rerun
from pyapp.config.config import Config


class DuplicateReportWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / 'source'
        self.source.mkdir()
        self.output = self.source / 'reports'
        self.patch = mock.patch.object(Config, 'appDataDir', str(self.root / 'state'))
        self.patch.start()
        self.api = API()

    def tearDown(self):
        self.api.task_shutdown()
        self.api.workflow_stop()
        self.patch.stop()
        self.temp.cleanup()

    def inputs(self):
        return {'directory': str(self.source), 'outputDir': str(self.output), 'extensions': [], 'recursive': True, 'limit': 100}

    def template(self):
        return self.api.workflow_create_from_template({'templateId': 'builtin-duplicate-report'})['workflow']

    def test_real_complete_scan_report_and_query_source_identity(self):
        first = self.source / 'first.txt'
        first.write_text('same', encoding='utf-8')
        (self.source / 'second.txt').write_text('same', encoding='utf-8')
        workflow = self.template()
        self.assertTrue(self.api.workflow_preflight({'steps': workflow['steps'], 'input': self.inputs()})['valid'])
        result = self.api.workflow_run({'id': workflow['id'], 'input': self.inputs()})
        self.assertEqual(result['run']['status'], 'success', result)
        steps = result['context']['steps']
        target = Path(steps['report']['file'])
        self.assertEqual(len(steps['scan']['outputPaths']), 2)
        self.assertEqual(steps['report']['outputPaths'], [str(target)])
        with closing(load_workbook(target, read_only=True)) as book:
            self.assertEqual(len(list(book['逐文件核对'].rows)), 3)
        service = self.api._services[WorkflowMixin]
        self.assertNotIn(str(first.resolve()), service._workflow_generated_paths)
        self.assertIn(str(target.resolve()), service._workflow_generated_paths)
        again = self.api.workflow_run({'id': workflow['id'], 'input': self.inputs()})
        self.assertEqual(again['context']['steps']['scan']['scanned'], 2)
        self.assertNotEqual(again['context']['steps']['report']['file'], str(target))

    def test_partial_stops_before_report_and_empty_complete_gets_report(self):
        workflow = self.template()
        empty = self.api.workflow_run({'id': workflow['id'], 'input': self.inputs()})
        self.assertEqual(empty['run']['status'], 'success', empty)
        count = len(list(self.output.glob('*.xlsx')))
        for index in range(101):
            (self.source / f'{index}.txt').write_text('same', encoding='utf-8')
        partial = self.api.workflow_run({'id': workflow['id'], 'input': self.inputs()})
        self.assertEqual(partial['run']['status'], 'partial', partial)
        self.assertEqual(len(partial['run']['steps']), 1)
        self.assertEqual(len(list(self.output.glob('*.xlsx'))), count)

    def test_catalog_declares_mode_groups_summary_and_report_dependencies(self):
        operations = {item['id']: item for item in self.api.operations_list()['operations']}
        self.assertTrue(operations['file_deduplicate_report']['workflow'])
        self.assertTrue(operations['file_deduplicate_report']['cancellable'])
        scan = operations['file_deduplicate']
        self.assertIn('summary', [field['path'] for field in scan['resultFields']])
        workflow = self.template()
        checked = self.api.workflow_preflight({'steps': workflow['steps'], 'input': {**self.inputs(), 'outputDir': ''}})
        self.assertFalse(checked['valid'])

    def test_partial_continue_resume_refreshes_query_report_and_transitive_archive(self):
        for name in ('a.txt', 'b.txt', 'c.txt'):
            (self.source / name).write_bytes(b'abc')
        workflow = self.template()
        workflow['steps'][0]['onPartial'] = 'continue'
        workflow['steps'].insert(1, {'id': 'independent', 'method': 'text_case_transform',
                                    'args': {'content': 'stable', 'mode': 'upper'}, 'onError': 'stop'})
        workflow['steps'].append({'id': 'archive', 'method': 'file_compress',
                                  'args': {'items': '{{steps.report.outputPaths}}', 'outputDir': '{{input.outputDir}}',
                                           'archiveName': 'snapshot', 'format': 'zip'}, 'onError': 'stop'})
        workflow = self.api.workflow_save(workflow)['workflow']
        def unavailable(path, expected, check):
            if path.name == 'c.txt':
                raise PermissionError('temporary read failure')
            return _stable_hash(path, expected, check)
        with mock.patch('api.core.file_duplicates._stable_hash', side_effect=unavailable):
            first = self.api.workflow_run({'id': workflow['id'], 'input': self.inputs()})
        self.assertEqual(first['run']['status'], 'partial', first)
        old_report = Path(first['context']['steps']['report']['file'])
        old_bytes = old_report.read_bytes()
        self.assertEqual(first['context']['steps']['scan']['totalGroups'], 1)
        (self.source / 'b.txt').write_bytes(b'different')
        (self.source / 'c.txt').write_bytes(b'now unique bytes')
        resumed = self.api.workflow_run({'id': workflow['id'], 'input': self.inputs(), '_resumeRunId': first['run']['id']})
        self.assertEqual(resumed['run']['status'], 'success', resumed)
        steps = resumed['context']['steps']
        self.assertEqual(steps['scan']['groups'], [])
        self.assertEqual(steps['scan']['outputAssets'], [])
        self.assertEqual(steps['scan']['outputPaths'], [])
        self.assertTrue(resumed['run']['steps'][1]['reused'])
        for step in resumed['run']['steps'][2:]:
            self.assertFalse(step.get('reused'))
            self.assertEqual(step['resumeInfo']['retainedOutputCount'], 0)
            self.assertIn('前序依赖', step['resumeInfo']['reason'])
        current_report = Path(steps['report']['file'])
        self.assertNotEqual(old_report, current_report)
        self.assertEqual(old_report.read_bytes(), old_bytes)
        with closing(load_workbook(current_report, read_only=True)) as book:
            self.assertEqual(len(list(book['逐文件核对'].rows)), 1)
            self.assertEqual(book['扫描说明']['B2'].value, '完整')
        with zipfile.ZipFile(steps['archive']['file']) as package:
            self.assertEqual(package.namelist(), [current_report.name])
            self.assertEqual(package.read(current_report.name), current_report.read_bytes())
        again = self.api.workflow_run({'id': workflow['id'], 'input': self.inputs(), '_resumeRunId': resumed['run']['id']})
        self.assertTrue(all(step['reused'] for step in again['run']['steps']))
        self.assertEqual(again['context']['steps']['report']['file'], str(current_report))

    def test_rerun_dependencies_use_actual_binding_syntax_and_nested_values(self):
        changed = {'scan'}
        self.assertTrue(_depends_on_rerun({'nested': ['prefix {{ steps.scan.summary }}']}, changed))
        self.assertTrue(_depends_on_rerun({'all': '{{steps}}'}, changed))
        self.assertFalse(_depends_on_rerun({'text': '{{input.content}}', 'other': '{{steps.unrelated.output}}'}, changed))
        self.assertFalse(_depends_on_rerun({'text': 'steps.scan.summary'}, changed))
        self.assertFalse(_depends_on_rerun({'text': '{{steps}}'}, set()))


if __name__ == '__main__':
    unittest.main()
