"""Real workbook processing through the desktop browser bridge."""

import hashlib
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect


def verify_excel_processing(api, page, root, chosen, report_dir):
    source = root / 'excel-source.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Records'
    sheet.append(['Report'])
    sheet.append(['code', 'name', 'date'])
    sheet.append(['002', ' Beta ', datetime(2025, 2, 1)])
    sheet.append(['001', ' Alpha ', datetime(2025, 1, 1)])
    sheet.append(['001', 'Alpha', datetime(2025, 1, 1)])
    book.save(source)
    book.close()
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    newer = root / 'excel-newer.xlsx'
    book = Workbook()
    book.active.append(['Report'])
    book.active.append(['new_code', 'new_value'])
    book.active.append(['B', 7])
    book.save(newer)
    book.close()
    bad = root / 'excel-corrupt.xlsx'
    bad.write_text('not a workbook', encoding='utf-8')
    navigation = page.get_by_role('navigation', name='工具', exact=True)
    navigation.get_by_role('button', name='Excel 工具', exact=True).click()
    tabs = page.get_by_role('navigation', name='子功能', exact=True)
    tabs.get_by_role('button', name='清洗与处理', exact=True).click()
    view = page.locator('.active-tool-view')
    panel = view.locator('.panel:visible').filter(has=page.get_by_role('heading', name='清洗 / 分组 / 排序 / 导出', exact=True))
    header = view.locator('.el-form-item').filter(has_text='表头所在行').locator('input')
    header.fill('2')
    header.press('Tab')
    chosen[:] = [source]
    panel.get_by_role('button', name='选择 / 接收主 Excel', exact=True).click()
    expect(panel.get_by_role('button', name='执行处理', exact=True)).to_be_enabled()

    def select_field(container, label, value):
        field = container.locator('.el-form-item').filter(has=page.locator('.el-form-item__label').filter(has_text=label))
        field.locator('.el-select__wrapper').click()
        choices = field.get_by_role('combobox').get_attribute('aria-controls')
        page.locator('[id="' + choices + '"]').get_by_role('option', name=value, exact=True).click()
        page.keyboard.press('Escape')

    view.get_by_text('去除文本首尾空格', exact=True).click()
    expect(view.get_by_role('checkbox', name='去除文本首尾空格', exact=True)).to_be_checked()
    select_field(panel, '去重字段', 'code')
    select_field(panel, '排序字段', 'date')
    panel.get_by_role('button', name='选择目录', exact=True).click()
    output_dir = str(root / 'outputs')
    expect(panel.get_by_placeholder('留空则自动创建', exact=True)).to_have_value(output_dir)
    panel.get_by_role('button', name='比较处理前后样本', exact=True).click()
    expect(panel.locator('.comparison-grid .el-table').nth(1).locator('.el-table__body tr')).to_have_count(2)
    after = panel.locator('.comparison-grid .el-table').nth(1).locator('.el-table__body tr')
    expect(after.first).to_contain_text('001')
    expect(after.nth(1)).to_contain_text('002')
    panel.get_by_role('button', name='执行处理', exact=True).click()
    expect(panel.get_by_text('处理摘要', exact=True)).to_be_visible(timeout=60000)
    exported = list((root / 'outputs').rglob('*.xlsx'))
    assert len(exported) == 1, exported
    result = load_workbook(exported[0])
    rows = list(result.active.values)
    result.close()
    assert rows[0] == ('code', 'name', 'date'), rows
    assert [row[:2] for row in rows[1:]] == [('001', 'Alpha'), ('002', 'Beta')], rows
    assert [row[2] for row in rows[1:]] == [datetime(2025, 1, 1), datetime(2025, 2, 1)], rows
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest

    page.set_viewport_size({'width': 980, 'height': 760})
    expect(page.locator('.el-message:visible')).to_have_count(0)
    for theme in ('light', 'dark'):
        page.evaluate("theme => { document.documentElement.dataset.theme = theme; document.documentElement.classList.toggle('dark', theme === 'dark') }", theme)
        panel.get_by_role('button', name='选择 / 接收主 Excel', exact=True).scroll_into_view_if_needed()
        page.screenshot(path=str(report_dir / f'excel-{theme}.png'), full_page=True)
        box = panel.get_by_role('button', name='选择 / 接收主 Excel', exact=True).bounding_box()
        assert box and box['x'] >= 0 and box['x'] + box['width'] <= 980, box
        panel.get_by_role('button', name='执行处理', exact=True).scroll_into_view_if_needed()
        box = panel.get_by_role('button', name='执行处理', exact=True).bounding_box()
        assert box and 0 <= box['y'] and box['y'] + box['height'] <= 760, box
        page.screenshot(path=str(report_dir / f'excel-{theme}-results.png'), full_page=True)
        tag = panel.locator('.result-list .output-path').filter(has_text='主表：')
        tag_box = tag.bounding_box()
        panel_box = panel.bounding_box()
        assert tag_box and panel_box and tag_box['x'] >= panel_box['x'] and tag_box['x'] + tag_box['width'] <= min(980, panel_box['x'] + panel_box['width']), 'Export path overflows the result panel'
    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.evaluate("document.documentElement.dataset.theme = 'light'; document.documentElement.classList.remove('dark')")
    chosen[:] = [newer]
    panel.get_by_role('button', name='选择 / 接收主 Excel', exact=True).click()
    expect(panel.get_by_text('处理摘要', exact=True)).to_have_count(0)
    expect(panel.locator('.comparison-grid')).to_have_count(0)
    expect(panel.get_by_placeholder('留空则自动创建', exact=True)).to_have_value(output_dir)
    expect(panel.get_by_role('button', name='执行处理', exact=True)).to_be_enabled()
    chosen[:] = [bad]
    panel.get_by_role('button', name='选择 / 接收主 Excel', exact=True).click()
    expect(view.locator('.excel-error')).to_be_visible()
    expect(panel.get_by_role('button', name='执行处理', exact=True)).to_be_disabled()

    page.evaluate("""() => {
      const original = window.pywebview.api.excel_preview;
      window.__excelOriginal = original;
      window.__slowStarted = false;
      window.__slowFinished = false;
      window.pywebview.api.excel_preview = async (...args) => {
        const result = await original(...args);
        if (args[0].filePath.endsWith('excel-source.xlsx')) {
          window.__slowStarted = true;
          await new Promise(resolve => setTimeout(resolve, 1500));
          window.__slowFinished = true;
        }
        return result;
      };
    }""")
    chosen[:] = [source]
    panel.get_by_role('button', name='选择 / 接收主 Excel', exact=True).click()
    page.wait_for_function('window.__slowStarted')
    chosen[:] = [newer]
    panel.get_by_role('button', name='选择 / 接收主 Excel', exact=True).click()
    page.wait_for_function('window.__slowFinished')
    expect(panel.get_by_role('button', name='执行处理', exact=True)).to_be_enabled()
    select_field(panel, '排序字段', 'new_value')
    expect(panel.locator('.input-path')).to_contain_text(newer.name)
    page.evaluate('window.pywebview.api.excel_preview = window.__excelOriginal')

    tabs.get_by_role('button', name='数据质检', exact=True).click()
    profile = view.locator('.panel:visible').filter(has=page.get_by_role('heading', name='Excel 数据质检', exact=True))
    chosen[:] = [source]
    profile.get_by_role('button', name='选择 / 接收 Excel', exact=True).click()
    expect(profile.get_by_role('button', name='开始质检', exact=True)).to_be_enabled()
    profile.get_by_role('button', name='开始质检', exact=True).click()
    expect(profile.get_by_text('Records：3 行，3 列', exact=True)).to_be_visible(timeout=60000)
    profile.get_by_placeholder('报告输出目录', exact=True).fill(str(root / 'quality-output'))
    profile.get_by_role('button', name='导出质量报告', exact=True).click()
    expect(profile.get_by_role('button', name='打开报告', exact=True)).to_be_visible(timeout=60000)
    reports = list((root / 'quality-output').rglob('*'))
    assert any(path.is_file() and path.stat().st_size > 0 for path in reports), reports
    quality = load_workbook(next(path for path in reports if path.suffix == '.xlsx'))
    assert dict(list(quality['质量概览'].values)[1:])['数据行数'] == 3
    assert [row[0] for row in list(quality['字段画像'].values)[1:]] == ['code', 'name', 'date']
    quality.close()

    tabs.get_by_role('button', name='按列拆分', exact=True).click()
    split = view.locator('.panel:visible').filter(has=page.get_by_role('heading', name='按列拆分工作簿', exact=True))
    split.get_by_role('button', name='选择 / 接收 Excel', exact=True).click()
    expect(split.get_by_role('button', name='开始拆分', exact=True)).to_be_enabled()
    select_field(split, '拆分列', 'code')
    split.get_by_role('button', name='开始拆分', exact=True).click()
    expect(split.get_by_role('button', name='打开目录', exact=True)).to_be_visible(timeout=60000)
    expect(split.locator('.el-table__body tr')).to_have_count(2)
    split_dir = Path(split.get_by_placeholder('自动创建', exact=True).input_value())
    split_files = list(split_dir.glob('*.xlsx'))
    assert len(split_files) == 2, split_files
    split_rows = []
    for path in split_files:
        workbook = load_workbook(path)
        values = list(workbook.active.values)
        workbook.close()
        assert values[0] == ('code', 'name', 'date'), values
        assert len({row[0] for row in values[1:]}) == 1, values
        split_rows.extend(values[1:])
    assert len(split_rows) == 3 and any(row[1] == ' Alpha ' for row in split_rows), split_rows
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    report = {'passed': True, 'checks': ['header row 2', 'typed date sort', 'trim and deduplicate preview/export parity', 'source unchanged', 'file switch clears results and preserves directory', 'corrupt file error', 'slow A cannot overwrite B', 'quality retains original 3 rows', 'quality report generated', 'split into two groups', '980x760 light/dark action bounds']}
    (report_dir / 'excel-processing.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report
