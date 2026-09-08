"""Verify result routing through real views using isolated, imported history fixtures."""

from __future__ import annotations

import re
import shutil
import subprocess
import time

import fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image
from playwright.sync_api import expect


def verify_result_handoff(api, page, root, chosen, picker_calls, report_dir):
    sources = root / 'handoff-fixtures'
    sources.mkdir()
    files = {}
    for suffix in ('png', 'pdf', 'docx', 'xlsx'):
        files[suffix] = []
        for number in (1, 2):
            path = sources / f'relay-{number}.{suffix}'
            if suffix == 'png':
                Image.new('RGB', (80, 60), (number * 70, 90, 130)).save(path)
            elif suffix == 'pdf':
                with fitz.open() as pdf:
                    pdf.new_page().insert_text((30, 40), f'Handoff {number}')
                    pdf.save(path)
            elif suffix == 'docx':
                doc = Document()
                doc.add_paragraph(f'Handoff {number}')
                doc.save(path)
            else:
                book = Workbook()
                book.active.append(['name', 'value'])
                book.active.append(['handoff', number])
                book.save(path)
                book.close()
            files[suffix].append(path)
    note = sources / 'relay-note.txt'
    note.write_text('Local handoff verification', encoding='utf-8')
    unknown = sources / 'relay.unknown'
    unknown.write_bytes(b'local artifact')
    video = sources / 'relay.mp4'
    capabilities = api.capabilities_get()['capabilities']
    if capabilities.get('ffmpeg', {}).get('available'):
        subprocess.run([shutil.which('ffmpeg'), '-hide_banner', '-loglevel', 'error', '-f', 'lavfi', '-i',
                        'color=c=blue:s=64x64:d=1', '-c:v', 'mpeg4', '-y', str(video)], check=True,
                       capture_output=True, timeout=30,
                       creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    else:
        video.write_bytes(b'fixture for disabled dependency checks only')

    all_paths = [*files['png'], *files['pdf'], *files['docx'], *files['xlsx'], note, unknown, video]
    outputs = [{'path': str(path), 'name': path.name, 'kind': 'file'} for path in all_paths]
    outputs.extend([{'path': str(sources), 'name': 'source-directory', 'kind': 'directory'},
                    {'path': str(sources / 'missing.pdf'), 'name': 'missing.pdf', 'kind': 'file', 'exists': False}])
    imported = api.task_import_legacy({'tasks': [{
        'id': 'e2e-result-handoff', 'method': 'file_batch_copy', 'status': 'success',
        'createdAt': time.time(), 'message': '接力路由验证夹具', 'outputs': outputs,
    }]})
    assert imported['imported'] == 1, imported
    task_ids = {task['id'] for task in api.task_list({'limit': 1000})['tasks']}

    def open_results(paths=None):
        page.get_by_role('button', name='打开任务中心', exact=True).click()
        page.get_by_role('button', name='刷新', exact=True).click()
        card = page.locator('.task-card').filter(has_text='接力路由验证夹具')
        card.get_by_role('button', name='检查结果 / 继续处理', exact=False).click()
        dialog = page.get_by_role('dialog', name='处理结果', exact=True)
        expect(dialog).to_be_visible()
        if paths is not None:
            dialog.get_by_role('button', name='清空选择', exact=True).click()
            expect(dialog.get_by_role('button', name='交给下一工具', exact=True)).to_be_disabled()
            for path in paths:
                dialog.get_by_text(path.name, exact=True).click()
                expect(dialog.get_by_role('checkbox', name=path.name, exact=True)).to_be_checked()
        return dialog

    def choose_route(dialog, label):
        dialog.locator('.handoff-options .el-select__wrapper').click()
        option = page.get_by_role('option').filter(has_text=re.compile('^' + re.escape(label) + r' · \d+ 个文件'))
        expect(option).to_be_visible()
        return option

    page.evaluate("localStorage.setItem('ppx-v2-modules', JSON.stringify({video: true, document: false}))")
    page.reload()
    dialog = open_results([note])
    expect(choose_route(dialog, '建立文档索引')).to_be_disabled()
    page.keyboard.press('Escape')
    dialog.get_by_role('button', name='取消', exact=True).click()
    page.evaluate("localStorage.setItem('ppx-v2-modules', JSON.stringify({video: true, document: true}))")
    page.reload()

    # Mixed selection must describe the subset before the user sends it.
    dialog = open_results()
    expect(dialog.get_by_role('checkbox', name='source-directory', exact=True)).to_be_disabled()
    expect(dialog.get_by_role('checkbox', name='missing.pdf', exact=True)).to_be_disabled()
    choose_route(dialog, '压缩图片').click()
    expect(dialog.get_by_text('本次跳过', exact=False)).to_be_visible()
    expect(page.locator('.el-select-dropdown:visible')).to_have_count(0)
    expect(page.locator('.el-message:visible')).to_have_count(0)
    page.screenshot(path=str(report_dir / 'handoff-results-light.png'), full_page=True)
    page.evaluate("document.documentElement.dataset.theme = 'dark'; document.documentElement.classList.add('dark')")
    page.set_viewport_size({'width': 980, 'height': 760})
    footer = dialog.get_by_role('button', name='交给下一工具', exact=True).bounding_box()
    assert footer and footer['y'] + footer['height'] <= 760, 'Result actions overflow the short window'
    skipped_box = dialog.locator('.skip-note').bounding_box()
    assert skipped_box and skipped_box['y'] + skipped_box['height'] < footer['y'], 'Skipped files must be explained before sending'
    path_box = dialog.locator('.asset small').first.bounding_box()
    assert path_box and path_box['height'] >= 12, 'Output paths must remain readable inside the checkbox group'
    page.screenshot(path=str(report_dir / 'handoff-results-dark.png'), full_page=True)
    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.evaluate("document.documentElement.dataset.theme = 'light'; document.documentElement.classList.remove('dark')")
    dialog.get_by_role('button', name='取消', exact=True).click()

    routes = [
        ('image/compress', '压缩图片', files['png'], None, 'managed'),
        ('image/watermark', '添加图片水印', files['png'], None, 'managed'),
        ('image/rotate', '旋转图片', files['png'], None, 'managed'),
        ('pdf/merge', '合并 PDF', files['pdf'], None, 'managed'),
        ('word/merge', '合并 Word', files['docx'], None, 'managed'),
        ('conversion/images-pdf', '图片合成 PDF', files['png'], 'flyingmouse', '添加图片'),
        ('conversion/merge-pdf', '转换引擎合并 PDF', files['pdf'], 'flyingmouse', '添加 PDF'),
        ('conversion/universal', '转换格式', files['png'], 'flyingmouse', '添加文件'),
        ('pdf/compress', '压缩 PDF', files['pdf'][:1], None, '选择 PDF'),
        ('pdf/ocr', 'PDF OCR', files['pdf'][:1], 'ocr', 'ocr'),
        ('excel/merge', '合并 Excel', files['xlsx'], None, '选择文件'),
        ('excel/profile', 'Excel 数据质检', files['xlsx'][:1], None, '选择 / 接收 Excel'),
        ('excel/process', '清洗 Excel', files['xlsx'][:1], None, '选择 / 接收主 Excel'),
        ('excel/split', '拆分 Excel', files['xlsx'][:1], None, '选择 / 接收 Excel'),
        ('document/index', '建立文档索引', [note], None, '添加文件'),
        ('file/archive', '打包归档', [unknown, note], None, '添加文件'),
        ('video/compress', '压缩视频', [video], 'ffmpeg', '选择'),
        ('video/cut', '截取视频', [video], 'ffmpeg', '选择'),
    ]
    verified, unavailable = [], []
    for route_id, label, paths, dependency, primary in routes:
        dialog = open_results(paths)
        option = choose_route(dialog, label)
        if dependency and capabilities.get(dependency, {}).get('available') is False:
            expect(option).to_be_disabled()
            page.keyboard.press('Escape')
            dialog.get_by_role('button', name='取消', exact=True).click()
            unavailable.append(route_id)
            continue
        option.click()
        dialog.get_by_role('button', name='交给下一工具', exact=True).click()
        notice = page.locator('.incoming-notice')
        expect(notice).to_contain_text(f'待接收：{len(paths)} 个结果 → {label}')
        if route_id == 'image/watermark':
            # Returning to the same route after a manual tab change must work.
            page.get_by_role('navigation', name='子功能', exact=True).get_by_role('button', name='批量压缩', exact=True).click()
            notice.get_by_role('button', name='返回目标操作', exact=True).click()
            page.locator('.panel:visible').get_by_text('图片', exact=True).click()
            chosen[:] = [files['png'][0]]
            before = len(picker_calls)
            page.locator('.el-form-item:visible').filter(has_text='水印图片').get_by_role('button', name='选择', exact=True).click()
            assert len(picker_calls) == before + 1, 'Auxiliary picker was hijacked'
            expect(notice).to_be_visible()
        before = len(picker_calls)
        if primary == 'managed':
            page.get_by_role('button', name=f'使用上一步结果：{len(paths)} 个文件', exact=True).click()
        elif primary == 'ocr':
            page.locator('.ocr-panel:visible .source-card').click()
        else:
            page.locator('.active-tool-view button:visible').filter(has_text=re.compile('^' + re.escape(primary) + '$')).first.click()
        expect(notice).to_be_hidden()
        assert len(picker_calls) == before, f'{route_id} opened the native dialog instead of consuming results'
        content = page.locator('.active-tool-view').inner_text()
        values = page.locator('.active-tool-view input:visible').evaluate_all('(inputs) => inputs.map(input => input.value)')
        for path in paths:
            assert path.name in content or any(path.name in value for value in values), f'{route_id} lost {path.name}'
        verified.append(route_id)

    # Receiving the same archive results a second time must not duplicate its queue.
    dialog = open_results([unknown, note])
    choose_route(dialog, '打包归档').click()
    dialog.get_by_role('button', name='交给下一工具', exact=True).click()
    page.locator('.active-tool-view button:visible').filter(has_text=re.compile('^添加文件$')).first.click()
    archive_rows = page.locator('.active-tool-view .el-table:visible .el-table__body tr')
    expect(archive_rows).to_have_count(2)
    page.screenshot(path=str(report_dir / 'handoff-received.png'), full_page=True)
    assert {task['id'] for task in api.task_list({'limit': 1000})['tasks']} == task_ids, 'Receiving results must not execute tasks'
    return {'verifiedRoutes': verified, 'unavailableDependencies': unavailable,
            'checks': ['mixed selection', 'single-file limits', 'disabled module', 'missing assets',
                       'auxiliary picker isolation', 'same-route return', 'deduplicated append']}
