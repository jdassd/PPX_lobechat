import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from PIL import Image

from api.api import API
from api.operations import OPERATIONS
from api.utils.error_handler import api_success
from api.workflow import BUILTIN_WORKFLOWS, WorkflowMixin
from pyapp.config.config import Config


class WorkflowApi(WorkflowMixin):
    def image_batch_compress(self, options=None):
        options = options or {}
        files = options.get('files') or []
        made = [path + '.small' for path in files if Path(path).is_file()]
        return api_success('compressed', outputPaths=made, partial=len(made) != len(files))

    def file_compress(self, options=None):
        return api_success('archived', file=str(Path((options or {})['outputDir']) / 'archive.zip'))

    def text_case_transform(self, options=None):
        return api_success('ok', output=str((options or {}).get('value', '')).upper())


class WorkflowTemplateTests(unittest.TestCase):
    def test_templates_are_bounded_to_real_contracts(self):
        templates = {item['id']: item for item in BUILTIN_WORKFLOWS}
        image = templates['builtin-image-archive']
        self.assertEqual(image['inputExample']['archiveName'], '图片分享包')
        self.assertEqual(image['steps'][0]['args']['quality'], 82)
        self.assertEqual(image['steps'][1]['args']['items'], '{{steps.compress.outputPaths}}')
        self.assertNotIn('password', image['steps'][1]['args'])
        excel = templates['builtin-excel-quality']
        self.assertEqual(excel['steps'][0]['args']['exportCombined'], True)
        self.assertEqual(excel['steps'][0]['args']['exportGroups'], False)
        self.assertEqual(excel['steps'][1]['args']['filePath'], '{{steps.clean.combinedPath}}')
        self.assertEqual(excel['steps'][1]['args']['headerRow'], 1)
        for method in ('excel_process', 'excel_merge_tables'):
            names = {field['name'] for field in OPERATIONS[method].fields}
            self.assertTrue({'trimText', 'deduplicateColumns'} <= names)

    def test_real_pngs_are_compressed_then_archived_without_changing_sources(self):
        with tempfile.TemporaryDirectory() as directory, mock.patch.object(Config, 'appDataDir', str(Path(directory) / 'state')):
            root = Path(directory)
            source = root / 'source.png'
            Image.new('RGB', (24, 16), 'red').save(source)
            original = source.read_bytes()
            api = API()
            created = api.workflow_create_from_template({'templateId': 'builtin-image-archive'})
            result = api.workflow_run({'id': created['workflow']['id'], 'input': {
                'files': [str(source)], 'outputDir': str(root / 'output'), 'archiveName': 'share'}})
            api.task_shutdown()
            api.workflow_stop()
            self.assertEqual(result['code'], 0, result)
            archive = Path(result['context']['steps']['archive']['file'])
            self.assertTrue(archive.is_file())
            with zipfile.ZipFile(archive) as package:
                self.assertEqual(package.namelist(), ['source_compress.png'])
                self.assertEqual(package.read(package.namelist()[0]), Path(result['context']['steps']['compress']['outputPaths'][0]).read_bytes())
            self.assertEqual(source.read_bytes(), original)

    def test_real_excel_cleaning_uses_original_header_then_reports_clean_copy_header_one(self):
        with tempfile.TemporaryDirectory() as directory, mock.patch.object(Config, 'appDataDir', str(Path(directory) / 'state')):
            root = Path(directory)
            source = root / 'source.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(['说明'])
            sheet.append(['姓名', '城市'])
            sheet.append([' Alice ', ' 北京 '])
            sheet.append(['Alice', '北京'])
            book.save(source)
            original = source.read_bytes()
            api = API()
            created = api.workflow_create_from_template({'templateId': 'builtin-excel-quality'})
            result = api.workflow_run({'id': created['workflow']['id'], 'input': {
                'filePath': str(source), 'outputDir': str(root / 'output'), 'headerRow': 2,
                'trimText': True, 'deduplicateColumns': ['姓名']}})
            api.task_shutdown()
            api.workflow_stop()
            self.assertEqual(result['code'], 0, result)
            clean = Path(result['context']['steps']['clean']['combinedPath'])
            cleaned = load_workbook(clean).active
            self.assertEqual(cleaned.max_row, 2)
            self.assertEqual([cell.value for cell in cleaned[1]], ['姓名', '城市'])
            self.assertEqual([cell.value for cell in cleaned[2]], ['Alice', '北京'])
            report = load_workbook(result['context']['steps']['report']['output'])
            self.assertIn('质量概览', report.sheetnames)
            self.assertEqual(source.read_bytes(), original)

    def test_real_partial_resume_keeps_good_outputs_then_archives_all_inputs(self):
        with tempfile.TemporaryDirectory() as directory, mock.patch.object(Config, 'appDataDir', str(Path(directory) / 'state')):
            root = Path(directory)
            good, missing = root / 'good.png', root / 'later.png'
            Image.new('RGB', (12, 8), 'red').save(good)
            api = API()
            self.addCleanup(api.task_shutdown)
            self.addCleanup(api.workflow_stop)
            saved = api.workflow_create_from_template({'templateId': 'builtin-image-archive'})['workflow']
            inputs = {'files': [str(good), str(missing)], 'outputDir': str(root / 'output'), 'archiveName': 'complete'}
            first = api.workflow_run({'id': saved['id'], 'input': inputs})
            self.assertTrue(first['partial'], first)
            self.assertEqual(len(first['run']['steps']), 1)
            retained = Path(first['context']['steps']['compress']['outputPaths'][0])
            before = retained.stat().st_mtime_ns
            self.assertTrue(retained.is_file())
            self.assertFalse(list((root / 'output').glob('*.zip')))
            Image.new('RGB', (12, 8), 'blue').save(missing)
            resumed = api.workflow_run({'id': saved['id'], 'input': inputs, '_resumeRunId': first['run']['id']})
            self.assertEqual(resumed['run']['status'], 'success', resumed)
            self.assertEqual(retained.stat().st_mtime_ns, before)
            self.assertEqual(len(resumed['context']['steps']['compress']['outputPaths']), 2)
            with zipfile.ZipFile(resumed['context']['steps']['archive']['file']) as package:
                self.assertEqual(set(package.namelist()), {'good_compress.png', 'later_compress.png'})

    def test_partial_stop_preserves_outputs_and_skips_archive(self):
        with tempfile.TemporaryDirectory() as directory, mock.patch.object(Config, 'appDataDir', directory):
            root = Path(directory)
            good = root / 'good.png'
            good.write_bytes(b'png')
            api = WorkflowApi()
            saved = api.workflow_create_from_template({'templateId': 'builtin-image-archive'})
            run = api.workflow_run({'id': saved['workflow']['id'], 'input': {
                'files': [str(good), str(root / 'missing.png')], 'outputDir': str(root), 'archiveName': 'x'}})
            self.assertEqual(run['code'], 0)
            self.assertTrue(run['partial'])
            self.assertEqual(len(run['run']['steps']), 1)
            self.assertEqual(run['run']['steps'][0]['status'], 'partial')

    def test_default_partial_continues_for_existing_workflows(self):
        with tempfile.TemporaryDirectory() as directory, mock.patch.object(Config, 'appDataDir', directory):
            root = Path(directory)
            good = root / 'good.png'
            good.write_bytes(b'png')
            api = WorkflowApi()
            saved = api.workflow_save({'name': 'legacy', 'steps': [
                {'id': 'one', 'method': 'image_batch_compress', 'args': {'files': [str(good), str(root / 'missing.png')]}},
                {'id': 'two', 'method': 'text_case_transform', 'args': {'value': 'ok'}},
            ]})
            run = api.workflow_run({'id': saved['workflow']['id']})
            self.assertEqual(run['run']['steps'][1]['status'], 'success')


if __name__ == '__main__':
    unittest.main()
