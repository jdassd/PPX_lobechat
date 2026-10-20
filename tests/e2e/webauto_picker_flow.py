"""Exercise browser collection through the visible picker and export flow.

The collector's own Chromium is launched headless with a private CDP port only
inside this test. An independent Playwright client performs real DOM clicks on
the injected picker toolbar and fixture page; production browser options are
not changed.
"""

from __future__ import annotations

import json
import queue
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from playwright.sync_api import expect, sync_playwright


class _FixtureHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path != '/list':
            self.send_error(404)
            return
        body = b'''<!doctype html><html><head><meta charset="utf-8"><title>Fixture list</title></head>
          <body><main>
            <article class="result"><h2 class="title">Alpha result</h2><a href="/alpha">Open Alpha</a></article>
            <article class="result"><h2 class="title">Beta result</h2><a href="/beta">Open Beta</a></article>
          </main></body></html>'''
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *_args):
        pass


def _wait_for_cdp(profile: Path, timeout: float = 25) -> str:
    endpoint_file = profile / 'DevToolsActivePort'
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            port = endpoint_file.read_text(encoding='utf-8').splitlines()[0]
            return f'http://127.0.0.1:{int(port)}'
        except (OSError, ValueError, IndexError):
            time.sleep(0.1)
    raise AssertionError(f'Test Chromium did not publish its CDP endpoint: {endpoint_file}')


def verify_webauto_picker(api, page, root: Path, report_dir: Path):
    """Run environment → picker clicks → preview → collection → XLSX reopen."""
    import playwright.sync_api as playwright_api
    from playwright.sync_api._generated import BrowserType

    from pyapp.config.config import Config

    report_dir.mkdir(parents=True, exist_ok=True)
    fixture = ThreadingHTTPServer(('127.0.0.1', 0), _FixtureHandler)
    fixture_thread = threading.Thread(target=fixture.serve_forever, daemon=True)
    fixture_thread.start()
    profile = Path(Config.appDataDir) / 'webauto' / 'browser-profile'
    original_sync_playwright = playwright_api.sync_playwright
    original_launch_persistent_context = BrowserType.launch_persistent_context
    picker_commands = queue.Queue()
    picker_results = queue.Queue()
    picker_thread = None

    def test_launch_persistent_context(self, user_data_dir, *args, **kwargs):
        # Test-only options: isolate the profile and expose an ephemeral CDP
        # endpoint so the independent test client can operate the real UI.
        kwargs['headless'] = True
        launch_args = list(kwargs.get('args') or [])
        launch_args.append('--remote-debugging-port=0')
        kwargs['args'] = launch_args
        return original_launch_persistent_context(self, user_data_dir, *args, **kwargs)

    try:
        # Patch only the browser launch call. Keeping the original Playwright
        # objects intact preserves its expose_binding callback machinery.
        BrowserType.launch_persistent_context = test_launch_persistent_context
        navigation = page.get_by_role('navigation', name='工具')
        navigation.get_by_role('button', name='网页数据采集', exact=True).click()
        expect(page.get_by_text('网页自动化采集', exact=True)).to_be_visible()
        status = api.webauto_status()
        assert status.get('code') == 0 and status.get('ready'), status
        assert status.get('installed'), f'CI fixture requires Chromium to be installed: {status}'
        expect(page.get_by_role('button', name='打开浏览器开始点选', exact=True)).to_be_enabled()

        fixture_url = f'http://127.0.0.1:{fixture.server_port}/list'
        page.get_by_placeholder('例如：https://example.com/list').fill(fixture_url)
        page.get_by_role('button', name='打开浏览器开始点选', exact=True).click()
        endpoint = _wait_for_cdp(profile)

        def picker_client():
            try:
                with original_sync_playwright() as cdp_driver:
                    browser = cdp_driver.chromium.connect_over_cdp(endpoint)
                    try:
                        context = browser.contexts[0]
                        deadline = time.monotonic() + 30
                        pages = context.pages
                        picker_page = next((candidate for candidate in pages if candidate.url.startswith(fixture_url)), None)
                        if picker_page is None:
                            picker_page = next((candidate for candidate in pages if candidate.url == 'about:blank'), None)
                        assert picker_page is not None, f'Picker page not found: {[p.url for p in pages]}'
                        if not picker_page.url.startswith(fixture_url):
                            remaining_ms = max(1, int((deadline - time.monotonic()) * 1000))
                            picker_page.wait_for_url(lambda url: url.startswith(fixture_url), timeout=remaining_ms)
                        toolbar = picker_page.locator('#__ppx_bar')
                        remaining_ms = max(1, int((deadline - time.monotonic()) * 1000))
                        toolbar.wait_for(state='visible', timeout=remaining_ms)
                        picker_results.put(('ready', None))
                        while True:
                            command = picker_commands.get()
                            if command == 'stop':
                                break
                            if command == 'container':
                                picker_page.locator('.result .title').first.click()
                            elif command == 'field-mode':
                                toolbar.get_by_role('button', name='②选字段').click()
                            elif command == 'title-field':
                                picker_page.locator('.result .title').first.click()
                            elif command == 'link-field':
                                picker_page.locator('.result a').first.click()
                            elif command == 'done':
                                toolbar.get_by_role('button', name='✔ 完成选取').click()
                            elif command == 'screenshot':
                                picker_page.screenshot(path=str(report_dir / 'webauto-picker-browser.png'), full_page=True)
                            picker_state = picker_page.evaluate('''() => ({
                              mode: window.__ppxState?.mode,
                              container: window.__ppxState?.container,
                              fields: window.__ppxState?.fields?.map(({name, selector, attr}) => ({name, selector, attr})) || []
                            })''')
                            picker_results.put((command, picker_state))
                    finally:
                        browser.close()
            except BaseException as exc:
                picker_results.put(('error', exc))

        def picker_action(command):
            picker_commands.put(command)
            result_name, error = picker_results.get(timeout=30)
            if result_name == 'error':
                raise error
            assert result_name == command, (result_name, command, error)
            return error

        picker_thread = threading.Thread(target=picker_client, name='webauto-picker-cdp-test', daemon=True)
        picker_thread.start()
        ready_name, ready_error = picker_results.get(timeout=25)
        if ready_name == 'error':
            raise ready_error
        assert ready_name == 'ready', (ready_name, ready_error)

        # Real clicks on a repeated row, then on two fields through the
        # injected toolbar. No server picker state is seeded in this test.
        picker_browser_state = picker_action('container')
        assert picker_browser_state.get('container'), f'Picker did not accept container click in its UI: {picker_browser_state}'
        try:
            expect(page.get_by_text('共找到 2 条相似内容', exact=False)).to_be_visible(timeout=10000)
        except AssertionError as exc:
            raise AssertionError(f'Container click was not reflected in UI; browser={picker_browser_state}; api={api.webauto_pick_state()}') from exc
        picker_action('field-mode')
        picker_action('title-field')
        picker_action('link-field')
        expect(page.get_by_text('field1（文本内容）', exact=False)).to_be_visible(timeout=10000)
        expect(page.get_by_text('field2（链接地址）', exact=False)).to_be_visible(timeout=10000)

        picker_action('done')
        expect(page.get_by_text('第二步：核对要采集的内容', exact=True)).to_be_visible(timeout=15000)
        names = page.locator('.wa-step:visible .wa-card:nth-of-type(2) .el-table__body-wrapper td:first-child input')
        expect(names).to_have_count(2)
        names.nth(0).fill('标题')
        names.nth(1).fill('链接')

        page.get_by_role('button', name='读取样本并检查', exact=True).click()
        expect(page.get_by_text('已读取当前页的前 10 条样本，请核对字段', exact=True)).to_be_visible(timeout=20000)
        sample_rows = page.locator('.wa-step:visible > .el-table .el-table__body-wrapper tr')
        expect(sample_rows).to_have_count(2)
        page.get_by_role('button', name='下一步：正式采集', exact=True).click()
        page.get_by_role('button', name='开始采集', exact=True).click()
        completion_status = page.locator('.wa-step:visible .el-tag__content:text-is("已完成")')
        expect(completion_status).to_have_count(1, timeout=30000)
        expect(completion_status).to_be_visible()
        expect(page.locator('.wa-step:visible .el-table__body-wrapper tr')).to_have_count(2)
        page.get_by_role('button', name='检查完成，导出结果', exact=True).click()
        page.get_by_role('button', name='选择目录', exact=True).click()
        page.get_by_role('button', name='导出结果', exact=True).click()
        expect(page.get_by_text('结果已导出', exact=True)).to_be_visible(timeout=15000)
        picker_action('screenshot')

        output_path = root / 'outputs' / '采集结果.xlsx'
        assert output_path.is_file(), f'Exported XLSX missing: {output_path}'
        from openpyxl import load_workbook

        workbook = load_workbook(output_path, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
        assert rows == [('标题', '链接'), ('Alpha result', '/alpha'), ('Beta result', '/beta')], rows
        page.screenshot(path=str(report_dir / 'webauto-picker-flow.png'), full_page=True)
        result = {
            'passed': True,
            'environment': {'ready': status['ready'], 'installed': status['installed']},
            'url': fixture_url,
            'picker': {'containerRows': 2, 'fields': ['标题', '链接'], 'selectedViaToolbar': True},
            'previewRows': 2,
            'collectedRows': 2,
            'xlsxRows': rows,
            'output': str(output_path),
            'testOnlyBrowserOptions': ['headless=True', '--remote-debugging-port=0', 'isolated appDataDir profile'],
        }
        (report_dir / 'webauto-picker-flow.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        return result
    finally:
        if picker_thread is not None and picker_thread.is_alive():
            picker_commands.put('stop')
            picker_thread.join(timeout=8)
        try:
            api.webauto_pick_cancel()
        except Exception:
            pass
        BrowserType.launch_persistent_context = original_launch_persistent_context
        fixture.shutdown()
        fixture.server_close()
        fixture_thread.join(timeout=2)
