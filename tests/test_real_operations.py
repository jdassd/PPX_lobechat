"""Real file-processing regressions for the main document and image workflows."""

import json
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import fitz
from docx import Document
from PIL import Image

from api.api import API
from api.image import ImageTool
from api.pdf import PDF
from api.video import VideoTool
from api.word import WordTool
from pyapp.config.config import Config


class RealOperationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="ppx-real-operations-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        code_dir = mock.patch.object(Config, "codeDir", str(Path(__file__).resolve().parents[1]))
        code_dir.start()
        self.addCleanup(code_dir.stop)
        app_data = mock.patch.object(Config, "appDataDir", str(self.root / "app-data"))
        app_data.start()
        self.addCleanup(app_data.stop)

    def make_pdf(self, name, pages):
        path = self.root / name
        document = fitz.open()
        for text in pages:
            page = document.new_page(width=240, height=180)
            page.insert_text((20, 40), text)
        document.save(path)
        document.close()
        return path

    def page_texts(self, path):
        with fitz.open(path) as document:
            return [page.get_text().strip() for page in document]

    def test_pdf_merge_respects_file_and_page_selection_order(self):
        first = self.make_pdf("first.pdf", ["first-1", "first-2"])
        second = self.make_pdf("second.pdf", ["second-1", "second-2"])
        result = PDF().pdf_merge({
            "files": [
                {"path": str(first), "pageSpec": "2,1"},
                {"path": str(second), "pageSpec": "2"},
            ],
            "outputPath": str(self.root / "merged.pdf"),
        })
        self.assertEqual(result["code"], 0, result)
        self.assertEqual(result["mergedPages"], 3)
        self.assertEqual(self.page_texts(result["output"]), ["first-2", "first-1", "second-2"])

    def test_pdf_split_and_cut_create_expected_pages_and_reject_reversed_range(self):
        source = self.make_pdf("source.pdf", ["page-1", "page-2", "page-3"])
        split = PDF().pdf_split({
            "filePath": str(source), "pagesPerFile": 2, "outputDir": str(self.root / "split")
        })
        self.assertEqual(split["code"], 0, split)
        self.assertEqual([len(self.page_texts(path)) for path in split["files"]], [2, 1])
        self.assertEqual(self.page_texts(split["files"][0]), ["page-1", "page-2"])
        cut = PDF().pdf_cut({
            "filePath": str(source), "mode": "custom", "pageSpec": "3,1",
            "outputPath": str(self.root / "cut.pdf"),
        })
        self.assertEqual(cut["code"], 0, cut)
        self.assertEqual(self.page_texts(cut["output"]), ["page-3", "page-1"])
        invalid = PDF().pdf_cut({
            "filePath": str(source), "startPage": 3, "endPage": 1,
            "outputPath": str(self.root / "invalid.pdf"),
        })
        self.assertNotEqual(invalid["code"], 0, invalid)
        self.assertIn("开始页", invalid["msg"])
        self.assertFalse((self.root / "invalid.pdf").exists())

    def test_pdf_compression_keeps_pages_and_text_extraction_keeps_page_numbers(self):
        source = self.make_pdf("content.pdf", ["alpha marker", "beta marker"])
        api = PDF()
        compressed = api.pdf_compress({
            "filePath": str(source), "mode": "custom", "customDpi": 96,
            "outputPath": str(self.root / "compressed.pdf"),
        })
        self.assertEqual(compressed["code"], 0, compressed)
        self.assertEqual(self.page_texts(compressed["output"]), ["", ""])
        with fitz.open(compressed["output"]) as result_pdf:
            self.assertEqual(result_pdf.page_count, 2)
            self.assertTrue(all(page.get_images() for page in result_pdf))

        extracted = api.pdf_extract_text({
            "filePath": str(source), "pageSpec": "2", "saveFile": True,
            "outputDir": str(self.root / "text"),
        })
        self.assertEqual(extracted["code"], 0, extracted)
        self.assertEqual(extracted["pages"], 1)
        self.assertEqual(extracted["segments"][0]["page"], 2)
        self.assertIn("beta marker", extracted["preview"])
        self.assertNotIn("alpha marker", extracted["preview"])
        self.assertIn("beta marker", Path(extracted["output"]).read_text(encoding="utf-8"))

    def test_pdf_extract_images_exports_embedded_raster_and_filters_by_size(self):
        source = self.root / "with-image.pdf"
        image_path = self.root / "embedded.png"
        Image.new("RGB", (32, 24), (210, 30, 80)).save(image_path)
        document = fitz.open()
        page = document.new_page(width=240, height=180)
        page.insert_image(fitz.Rect(10, 10, 100, 80), filename=str(image_path))
        document.save(source)
        document.close()

        api = PDF()
        extracted = api.pdf_extract_images({
            "filePath": str(source), "outputDir": str(self.root / "images")
        })
        self.assertEqual(extracted["code"], 0, extracted)
        self.assertEqual(extracted["count"], 1)
        with Image.open(extracted["files"][0]) as output:
            self.assertEqual(output.size, (32, 24))
            self.assertEqual(output.convert("RGB").getpixel((0, 0)), (210, 30, 80))
        filtered = api.pdf_extract_images({
            "filePath": str(source), "minWidth": 33, "outputDir": str(self.root / "filtered")
        })
        self.assertEqual(filtered["code"], 0, filtered)
        self.assertEqual(filtered["count"], 0)

    def test_word_heading_and_paragraph_splitting_preserve_real_content(self):
        source = self.root / "chapters.docx"
        document = Document()
        document.add_heading("Chapter One", level=1)
        document.add_paragraph("one body")
        document.add_heading("Chapter Two", level=1)
        document.add_paragraph("two body")
        document.save(source)

        tool = WordTool()
        heading = tool.word_split({
            "filePath": str(source), "mode": "heading", "headingLevel": 1,
            "outputDir": str(self.root / "by-heading"),
        })
        self.assertEqual(heading["code"], 0, heading)
        self.assertEqual(len(heading["files"]), 2)
        heading_text = [[p.text for p in Document(path).paragraphs] for path in heading["files"]]
        self.assertEqual(heading_text[0][:2], ["Chapter One", "one body"])
        self.assertEqual(heading_text[1][:2], ["Chapter Two", "two body"])

        paragraphs = tool.word_split({
            "filePath": str(source), "mode": "paragraphs", "paragraphsPerFile": 2,
            "outputDir": str(self.root / "by-paragraph"),
        })
        self.assertEqual(paragraphs["code"], 0, paragraphs)
        self.assertEqual(len(paragraphs["files"]), 2)
        self.assertEqual(
            [[p.text for p in Document(path).paragraphs] for path in paragraphs["files"]],
            [["Chapter One", "one body"], ["Chapter Two", "two body"]],
        )
        missing_heading = tool.word_split({
            "filePath": str(source), "mode": "heading", "headingLevel": 3,
            "outputDir": str(self.root / "no-heading"),
        })
        self.assertNotEqual(missing_heading["code"], 0, missing_heading)
        self.assertIn("未检测到", missing_heading["msg"])

    def test_image_format_conversion_and_compression_produce_readable_outputs(self):
        source = self.root / "photo.png"
        Image.new("RGB", (80, 60), (40, 110, 190)).save(source)
        tool = ImageTool()
        converted = tool.image_format_convert({
            "files": [str(source)], "targetFormat": "jpg", "quality": 85,
            "outputDir": str(self.root / "converted"),
        })
        self.assertEqual(converted["code"], 0, converted)
        with Image.open(converted["files"][0]) as result:
            self.assertEqual(result.format, "JPEG")
            self.assertEqual(result.size, (80, 60))
        compressed = tool.image_batch_compress({
            "files": [str(source)], "quality": 70, "outputDir": str(self.root / "compressed")
        })
        self.assertEqual(compressed["code"], 0, compressed)
        self.assertEqual(len(compressed["items"]), 1)
        self.assertTrue(Path(compressed["items"][0]["output"]).is_file())
        with Image.open(compressed["items"][0]["output"]) as result:
            self.assertEqual(result.size, (80, 60))

    def test_image_watermark_crop_and_concat_change_pixels_geometry_and_order(self):
        first = self.root / "red.png"
        second = self.root / "blue.png"
        Image.new("RGB", (20, 16), (220, 20, 20)).save(first)
        Image.new("RGB", (12, 10), (20, 20, 220)).save(second)
        tool = ImageTool()

        watermarked = tool.image_add_watermark({
            "files": [str(first)], "text": "MARK", "fontSize": 12,
            "position": "top-left", "outputDir": str(self.root / "watermarked"),
        })
        self.assertEqual(watermarked["code"], 0, watermarked)
        with Image.open(watermarked["files"][0]) as result:
            self.assertEqual(result.size, (20, 16))
            with Image.open(first) as original:
                self.assertNotEqual(list(result.convert("RGB").getdata()), list(original.getdata()))

        cropped = tool.image_crop({
            "file": str(first), "x": 4, "y": 3, "width": 9, "height": 7,
            "outputDir": str(self.root / "cropped"),
        })
        self.assertEqual(cropped["code"], 0, cropped)
        with Image.open(cropped["file"]) as result:
            self.assertEqual(result.size, (9, 7))

        joined = tool.image_concat({
            "files": [str(first), str(second)], "direction": "horizontal", "spacing": 2,
            "outputDir": str(self.root / "joined"),
        })
        self.assertEqual(joined["code"], 0, joined)
        with Image.open(joined["file"]) as result:
            self.assertEqual(result.size, (34, 16))
            self.assertEqual(result.getpixel((2, 8)), (220, 20, 20, 255))
            self.assertEqual(result.getpixel((25, 8)), (20, 20, 220, 255))

    def test_format_center_conversion_image_pdf_and_merge_when_engine_is_available(self):
        api = API()
        capabilities = api.format_center_capabilities()
        if not capabilities.get("available"):
            self.skipTest(capabilities.get("detail") or "FlyingMouse Format 引擎未就绪")

        first_image = self.root / "one.png"
        second_image = self.root / "two.png"
        Image.new("RGB", (32, 24), (220, 40, 20)).save(first_image)
        Image.new("RGB", (32, 24), (20, 40, 220)).save(second_image)
        output_dir = self.root / "format-output"
        converted = api.format_center_images_to_pdf({
            "files": [str(first_image), str(second_image)], "outputDir": str(output_dir),
            "outputName": "images.pdf",
        })
        self.assertEqual(converted["code"], 0, converted)
        converted_pdf = converted.get("output") or converted.get("outputs", [{}])[0].get("path")
        self.assertTrue(converted_pdf and Path(converted_pdf).is_file(), converted)
        with fitz.open(converted_pdf) as result:
            self.assertEqual(result.page_count, 2)

        second_pdf = self.make_pdf("other.pdf", ["second input"])
        merged = api.format_center_merge_pdfs({
            "files": [str(converted_pdf), str(second_pdf)], "outputDir": str(output_dir),
            "outputName": "merged.pdf",
        })
        self.assertEqual(merged["code"], 0, merged)
        merged_pdf = merged.get("output") or merged.get("outputs", [{}])[0].get("path")
        self.assertTrue(merged_pdf and Path(merged_pdf).is_file(), merged)
        with fitz.open(merged_pdf) as result:
            self.assertEqual(result.page_count, 3)

        rasterized = api.format_center_convert({
            "files": [str(first_image)], "targetFormat": "jpg",
            "outputDir": str(output_dir / "raster"),
        })
        self.assertEqual(rasterized["code"], 0, rasterized)
        outputs = rasterized.get("outputs") or rasterized.get("files") or []
        self.assertTrue(outputs, rasterized)
        for item in outputs:
            output_path = item.get("path") if isinstance(item, dict) else item
            self.assertTrue(Path(output_path).is_file(), output_path)
            with Image.open(output_path) as image:
                self.assertEqual(image.format, "JPEG")
                self.assertEqual(image.size, (32, 24))

    @unittest.skipUnless(shutil.which('ffmpeg') and shutil.which('ffprobe'), 'FFmpeg 工具链未安装')
    def test_video_processing_preserves_streams_and_expected_duration(self):
        source = self.root / 'sample.mp4'
        subprocess.run([
            shutil.which('ffmpeg'), '-hide_banner', '-loglevel', 'error', '-y',
            '-f', 'lavfi', '-i', 'color=c=red:s=160x120:r=10',
            '-f', 'lavfi', '-i', 'sine=frequency=440:sample_rate=48000',
            '-t', '2', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', '-c:a', 'aac', str(source),
        ], check=True, capture_output=True, timeout=30,
            creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        second = self.root / 'second.mp4'
        shutil.copyfile(source, second)
        tool = VideoTool()
        common = {'filePath': str(source), 'outputDir': str(self.root / 'video')}
        operations = [
            (tool.video_compress, {'preset': 'small', 'ffPreset': 'ultrafast'}, 2, {'video', 'audio'}),
            (tool.video_cut, {'start': 0.5, 'end': 1.5}, 1, {'video', 'audio'}),
            (tool.video_extract_audio, {'audioFormat': 'wav'}, 2, {'audio'}),
            (tool.video_concat, {'files': [str(source), str(second)]}, 4, {'video', 'audio'}),
        ]
        for operation, options, duration, stream_types in operations:
            with self.subTest(operation=operation.__name__):
                result = operation({**common, **options})
                self.assertEqual(result['code'], 0, result)
                probe = subprocess.run([
                    shutil.which('ffprobe'), '-v', 'error', '-show_format', '-show_streams',
                    '-of', 'json', result['file'],
                ], check=True, capture_output=True, text=True, timeout=30,
                    creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
                info = json.loads(probe.stdout)
                self.assertEqual({stream['codec_type'] for stream in info['streams']}, stream_types)
                self.assertAlmostEqual(float(info['format']['duration']), duration, delta=0.3)
        invalid = tool.video_cut({**common, 'start': 3, 'end': 4})
        self.assertNotEqual(invalid['code'], 0)
        self.assertIn('超出视频时长', invalid['msg'])


if __name__ == "__main__":
    unittest.main()
