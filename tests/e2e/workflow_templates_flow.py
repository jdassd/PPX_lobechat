"""Browser evidence for the two bounded, reusable workflow templates."""

import hashlib
import json
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image
from playwright.sync_api import expect


def verify_workflow_templates(api, page, root, chosen, report_dir):
    fixtures = root / 'workflow-template-fixtures'
    fixtures.mkdir()
    image = fixtures / 'source.png'
    Image.new('RGB', (96, 64), 'navy').save(image)
    original_image = image.read_bytes()
    book_path = fixtures / 'source.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.append(['说明'])
    sheet.append(['姓名', '城市'])
    sheet.append([' Alice ', ' 北京 '])
    sheet.append(['Alice', '北京'])
    book.save(book_path)
    book.close()
    original_book = hashlib.sha256(book_path.read_bytes()).hexdigest()
    output_dir = fixtures / ('long-output-' * 2)
    output_dir.mkdir(parents=True)

    def wait_run(workflow_id, expected='success'):
        for _ in range(300):
            run = next((item for item in api.workflow_list()['runs'] if item['workflowId'] == workflow_id), None)
            if run and run['status'] != 'running':
                assert run['status'] == expected, run
                return run
            page.wait_for_timeout(100)
        raise AssertionError('workflow run timeout')

    def use_template(name):
        page.get_by_role('navigation', name='工具', exact=True).get_by_role('button', name='自动化工作流', exact=True).click()
        page.get_by_role('tab', name='工作流', exact=True).click()
        card = page.locator('.template-card').filter(has_text=name)
        expect(card).to_be_visible()
        card.get_by_role('button', name='使用模板', exact=True).click()
        editor = page.locator('.workflow-editor')
        expect(editor.get_by_role('button', name='立即运行', exact=True)).to_be_enabled()
        return editor

    def form_item(scope, label):
        return scope.locator('.operation-form > .el-form > .el-form-item').filter(has_text=label)

    def new_task(method, before):
        for _ in range(600):
            task = next((item for item in api.task_list()['tasks'] if item['id'] not in before and item['method'] == method), None)
            if task and task['status'] not in ('queued', 'running', 'canceling'):
                return task
            page.wait_for_timeout(100)
        raise AssertionError(f'No completed browser task: {method}')

    # Image archive: the browser selects the run input and edits the unsaved step
    # quality.  Immediate run must first persist that current edit.
    chosen[:] = [image]
    editor = use_template('图片压缩 → 归档')
    run_setup = editor.locator('.run-setup')
    form_item(run_setup, '输入文件').get_by_role('button', name='选择', exact=True).click()
    form_item(run_setup, '输出目录').locator('input').fill(str(output_dir))
    form_item(run_setup, '压缩包名称').locator('input').fill('browser-share')
    compress = editor.locator('.step-card').first
    archive_step = editor.locator('.step-card').nth(1)
    form_item(archive_step, '处理项').locator('.reference .el-select__wrapper').click()
    page.get_by_role('option', name='压缩分享副本', exact=True).click()
    expect(form_item(archive_step, '处理项').locator('input').first).to_have_value('{{steps.compress.outputPaths}}')
    compress.locator('.form-mode .el-switch').click()
    args = {
        'files': '{{input.files}}', 'outputDir': '{{input.outputDir}}',
        'mode': 'quality', 'quality': 63,
    }
    compress.locator('textarea').fill(json.dumps(args, ensure_ascii=False))
    workflow_id = next(item['id'] for item in api.workflow_list()['workflows'] if item['name'].startswith('图片压缩'))
    before_runs = len(api.workflow_list()['runs'])
    editor.get_by_role('button', name='立即运行', exact=True).click()
    expect(page.get_by_text('工作流执行完成', exact=True)).to_be_visible(timeout=60000)
    assert len(api.workflow_list()['runs']) == before_runs + 1
    image_run = wait_run(workflow_id)
    saved_image = next(item for item in api.workflow_list()['workflows'] if item['id'] == workflow_id)
    assert saved_image['steps'][0]['args']['quality'] == 63, saved_image['steps'][0]
    archive = Path(image_run['steps'][1]['result']['file'])
    compressed = Path(image_run['steps'][0]['result']['outputPaths'][0])
    assert archive.is_file() and compressed.is_file()
    with zipfile.ZipFile(archive) as package:
        assert package.namelist() == [compressed.name]
        assert package.read(compressed.name) == compressed.read_bytes()
    assert image.read_bytes() == original_image

    # Invalid JSON has a visible error and cannot create a run.
    page.get_by_role('tab', name='工作流', exact=True).click()
    compress = editor.locator('.step-card').first
    compress.locator('.form-mode').scroll_into_view_if_needed()
    advanced_switch = compress.locator('.form-mode input[role="switch"]')
    if advanced_switch.get_attribute('aria-checked') != 'true':
        compress.locator('.form-mode .el-switch').click()
    compress.locator('textarea:visible').fill('{bad')
    count = len(api.workflow_list()['runs'])
    editor.get_by_role('button', name='立即运行', exact=True).click()
    expect(page.get_by_text('JSON 尚未完整，请修正后保存', exact=True)).to_be_visible()
    assert len(api.workflow_list()['runs']) == count
    compress.locator('textarea:visible').fill(json.dumps(args, ensure_ascii=False))

    # A queue-originated workflow run gives ResultActions a real queue task id,
    # never the run id.  The history UI still exposes the per-step result action.
    queued_task = next(item for item in api.task_list()['tasks'] if item.get('workflowRunId') == image_run['id'])
    page.get_by_role('tab', name='运行记录', exact=True).click()
    history = page.locator('.run-list .el-collapse-item').filter(has_text='图片压缩 → 归档').first
    history.locator('.el-collapse-item__header').click()
    expect(history.get_by_text('打包压缩副本', exact=True)).to_be_visible()
    history.get_by_role('button', name='检查结果 / 继续处理', exact=False).first.click()
    result_dialog = page.get_by_role('dialog', name='处理结果', exact=True)
    expect(result_dialog).to_be_visible()
    result_dialog.locator('.handoff-options .el-select__wrapper').click()
    page.get_by_role('option').filter(has_text=re.compile(r'^旋转图片 · 1 个文件')).click()
    result_dialog.get_by_role('button', name='交给下一工具', exact=True).click()
    page.get_by_role('button', name='使用上一步结果：1 个文件', exact=True).click()
    previous = {item['id'] for item in api.task_list()['tasks']}
    page.locator('.panel:visible').get_by_role('button', name='开始处理', exact=True).click()
    child = new_task('image_rotate_flip', previous)
    assert child['status'] == 'success', child
    assert child['inputOrigins'][0]['sourceTaskId'] == queued_task['id'], child
    assert child['inputOrigins'][0]['sourceTaskId'] != image_run['id']
    assert Path(child['inputOrigins'][0]['inputPath']) == compressed

    # Excel template: run input remains file-level, while the second step consumes
    # the first step's combinedPath (not its generic output field).
    editor = use_template('Excel 清洗 → 质检报告')
    run_setup = editor.locator('.run-setup')
    chosen[:] = [book_path]
    form_item(run_setup, '输入文件').get_by_role('button', name='选择', exact=True).click()
    form_item(run_setup, '输出目录').locator('input').fill(str(output_dir))
    form_item(run_setup, '表头所在行').locator('input').fill('2')
    expect(form_item(run_setup, '清除文本首尾空格').locator('input[role="switch"]')).to_have_attribute('aria-checked', 'true')
    # The list field is a deliberate ordinary form interaction, not a hidden path assignment.
    form_item(run_setup, '去重字段').locator('textarea').fill('姓名')
    workflow_id = next(item['id'] for item in api.workflow_list()['workflows'] if item['name'].startswith('Excel 清洗'))
    editor.get_by_role('button', name='立即运行', exact=True).click()
    expect(page.get_by_text('工作流执行完成', exact=True)).to_be_visible(timeout=60000)
    excel_run = wait_run(workflow_id)
    clean = Path(excel_run['steps'][0]['result']['combinedPath'])
    report = Path(excel_run['steps'][1]['result']['output'])
    assert clean.is_file() and report.is_file()
    cleaned = load_workbook(clean)
    rows = list(cleaned.active.values)
    cleaned.close()
    assert rows == [('姓名', '城市'), ('Alice', '北京')], rows
    quality = load_workbook(report)
    assert '质量概览' in quality.sheetnames
    assert dict(list(quality['质量概览'].values)[1:])['数据行数'] == 1
    quality.close()
    assert hashlib.sha256(book_path.read_bytes()).hexdigest() == original_book

    # Partial completion must not archive a subset. The ordinary task-center retry
    # resumes this same workflow and retains the completed image bytes/timestamp.
    page.get_by_role('tab', name='工作流', exact=True).click()
    page.locator('.workflow-list-item').filter(has_text='图片压缩 → 归档').click()
    run_setup = editor.locator('.run-setup')
    missing = fixtures / 'restored.png'
    chosen[:] = [image, missing]
    form_item(run_setup, '输入文件').get_by_role('button', name='选择', exact=True).click()
    form_item(run_setup, '压缩包名称').locator('input').fill('partial-share')
    previous = {item['id'] for item in api.task_list()['tasks']}
    editor.get_by_role('button', name='立即运行', exact=True).click()
    expect(page.get_by_text('工作流部分完成，请查看步骤记录和已生成文件', exact=True)).to_be_visible(timeout=60000)
    partial = new_task('workflow_run', previous)
    assert partial['status'] == 'partial', partial
    assert len(partial['result']['run']['steps']) == 1
    assert not (output_dir / 'partial-share.zip').exists()
    retained = Path(partial['result']['context']['steps']['compress']['outputPaths'][0])
    retained_time = retained.stat().st_mtime_ns
    Image.new('RGB', (96, 64), 'green').save(missing)
    page.get_by_role('button', name='打开任务中心', exact=True).click()
    page.get_by_placeholder('搜索任务、输出路径或错误信息', exact=True).fill('workflow_run')
    page.get_by_role('button', name='刷新', exact=True).click()
    card = page.locator('.task-card').filter(has_text='部分成功').first
    expect(card.get_by_role('button', name='重试', exact=True)).to_be_enabled()
    previous = {item['id'] for item in api.task_list()['tasks']}
    card.get_by_role('button', name='重试', exact=True).click()
    retried = new_task('workflow_run', previous)
    assert retried['status'] == 'success' and retried['retryOf'] == partial['id'], retried
    assert retained.stat().st_mtime_ns == retained_time
    with zipfile.ZipFile(retried['result']['context']['steps']['archive']['file']) as package:
        assert len(package.namelist()) == 2

    page.get_by_role('navigation', name='工具', exact=True).get_by_role('button', name='自动化工作流', exact=True).click()
    page.get_by_role('tab', name='运行记录', exact=True).click()
    page.locator('.workflow-tool').get_by_role('button', name='刷新', exact=True).click()
    expect(page.locator('.workflow-tool > .el-loading-mask')).not_to_be_visible()
    latest = page.locator('.run-list .el-collapse-item').first
    latest.locator('.el-collapse-item__header').click()
    expect(latest.get_by_role('button', name='检查结果 / 继续处理', exact=False)).to_have_count(2)
    expect(latest.get_by_role('button', name='检查结果 / 继续处理', exact=False).last).to_be_visible()
    expect(latest.locator('.resume-info').first).to_contain_text(partial['workflowRunId'])
    expect(latest).to_contain_text('本次处理 1 项 / 沿用已完成输入 1 项 / 沿用输出 1 个')
    page.wait_for_timeout(350)  # Finish the Element Plus collapse animation before evidence.
    latest.scroll_into_view_if_needed()
    page.screenshot(path=str(report_dir / 'templates-results.png'))

    # At compact size the run input, paths, steps and history remain usable.
    page.set_viewport_size({'width': 980, 'height': 760})
    page.get_by_role('tab', name='工作流', exact=True).click()
    page.locator('.workflow-list-item').filter(has_text='图片压缩 → 归档').click()
    run_setup = editor.locator('.run-setup')
    for theme in ('light', 'dark'):
        page.evaluate("theme => { document.documentElement.dataset.theme = theme; document.documentElement.classList.toggle('dark', theme === 'dark') }", theme)
        page.wait_for_timeout(250)
        assert page.locator('.workflow-tool').evaluate('(node) => node.scrollWidth <= node.clientWidth')
        page.locator('.workflow-tool').evaluate('(node) => { node.scrollTop += node.querySelector(".run-setup").getBoundingClientRect().top - node.getBoundingClientRect().top - 12 }')
        bounds = editor.get_by_role('button', name='立即运行', exact=True).bounding_box()
        assert bounds and bounds['y'] >= 0 and bounds['y'] + bounds['height'] <= 760
        page.screenshot(path=str(report_dir / f'templates-{theme}.png'))
    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.evaluate("document.documentElement.dataset.theme = 'light'; document.documentElement.classList.remove('dark')")
    result = {
        'passed': True,
        'checks': ['template UI input and immediate save', 'paths field references outputPaths', 'image copy ZIP and source unchanged', 'invalid JSON blocks run', 'history per-step result handoff preserves queue task origin', 'Excel header row 2 clean copy then header row 1 report', 'partial completion stops archive and keeps outputs', 'task-center retry restores complete archive without rewriting success', '980x760 light/dark no horizontal overflow'],
        'runs': [image_run['id'], excel_run['id'], partial['workflowRunId'], retried['workflowRunId']],
    }
    (report_dir / 'templates.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return result
