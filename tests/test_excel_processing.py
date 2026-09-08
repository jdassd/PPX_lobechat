import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from api.excel import Excel


class ExcelProcessingTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.api = Excel()

    def workbook(self, name, rows):
        path = self.root / name
        book = Workbook()
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        return path

    def test_typed_cleaning_preview_and_export_share_rules(self):
        source = self.workbook('typed.xlsx', [
            ['Report'], ['id', 'name', 'rank', 'date', 'formula', '_tag'],
            ['001', ' Alice ', 2, datetime(2026, 8, 2), '=C3*2', ' x '],
            ['001', 'Alice', 2, datetime(2026, 8, 2), '=C4*2', 'x'],
            ['002', ' Bob ', 1, datetime(2026, 8, 1), '=C5*2', ' y '],
        ])
        options = {'filePath': str(source), 'headerRow': 2, 'trimText': True,
                   'deduplicateColumns': ['id', 'name'], 'sortBy': 'rank', 'formulaPolicy': 'preserve',
                   'exportCombined': True, 'exportGroups': False, 'exportJson': False,
                   'outputDir': str(self.root / 'out')}
        original = source.read_bytes()
        preview = self.api.excel_process_preview(options)
        self.assertEqual(preview['code'], 0, preview)
        self.assertFalse((self.root / 'out').exists())
        self.assertEqual([row['id'] for row in preview['after']], ['002', '001'])
        self.assertEqual([row['_tag'] for row in preview['after']], ['y', 'x'])
        self.assertEqual(preview['sampleRemovedDuplicates'], 1)
        self.assertEqual(preview['before'][0]['name'], ' Alice ')
        self.assertEqual(preview['after'][0]['formula'], '=C5*2')
        result = self.api.excel_process(options)
        self.assertEqual(result['code'], 0, result)
        self.assertEqual(result['summary']['sourceRows'], 3)
        self.assertEqual(result['summary']['totalRows'], 2)
        self.assertEqual(result['summary']['removedDuplicates'], 1)
        output = load_workbook(result['combinedPath'])
        try:
            sheet = output.active
            self.assertEqual(sheet['A2'].value, '002')
            self.assertEqual(sheet['A2'].data_type, 's')
            self.assertEqual(sheet['B2'].value, 'Bob')
            self.assertEqual(sheet['C2'].value, 1)
            self.assertIsInstance(sheet['D2'].value, datetime)
            self.assertEqual(sheet['E2'].value, '=C2*2')
            self.assertEqual(sheet['F2'].value, 'y')
        finally:
            output.close()
        self.assertEqual(source.read_bytes(), original)

    def test_sample_sort_preserves_date_rank_between_number_and_text(self):
        source = self.workbook('mixed.xlsx', [['id', 'mixed'], ['text', '0001'],
                                            ['date', datetime(2026, 8, 1)], ['number', 5]])
        options = {'filePath': str(source), 'sortBy': 'mixed', 'exportCombined': True,
                   'outputDir': str(self.root / 'out')}
        preview = self.api.excel_process_preview(options)
        result = self.api.excel_process(options)
        self.assertEqual(preview['code'], 0, preview)
        self.assertEqual(result['code'], 0, result)
        self.assertEqual([row['id'] for row in preview['after']], ['number', 'date', 'text'])
        self.assertEqual([row['id'] for row in result['sample']], ['number', 'date', 'text'])

    def test_preview_is_bounded_and_explicit_about_additional_tables(self):
        source = self.workbook('large.xlsx', [['value']] + [[index] for index in range(500)])
        result = self.api.excel_process_preview({'filePath': str(source), 'limit': 500,
                                               'mergeFiles': [{'path': str(self.root / 'other.xlsx')}]})
        self.assertEqual(result['code'], 0, result)
        self.assertEqual(result['sampleRows'], 100)
        self.assertEqual(len(result['before']), 100)
        self.assertTrue(result['sampleOnly'])
        self.assertFalse(result['includesMergeFiles'])
        self.assertIn('不包含附加分表', result['msg'])
        self.assertEqual(list(self.root.iterdir()), [source])

    def test_preview_errors_are_structured_and_missing_cache_is_explained(self):
        source = self.workbook('formula.xlsx', [['value', 'formula'], [1, '=A2*2']])
        invalid_options = [{'sortBy': 'absent'}, {'groupBy': 'absent'}, {'deduplicateColumns': ['absent']},
                           {'deduplicateColumns': 'value'}, {'sortOrder': 'sideways'}, {'formulaPolicy': 'values'}]
        for options in invalid_options:
            with self.subTest(options=options):
                result = self.api.excel_process_preview({'filePath': str(source), **options})
                self.assertNotEqual(result['code'], 0, result)
                self.assertTrue(result['msg'])
        cached = self.api.excel_process_preview({'filePath': str(source), 'formulaPolicy': 'values'})
        self.assertIn('缓存', cached['msg'])
        self.assertNotEqual(self.api.excel_process_preview([])['code'], 0)
        self.assertNotEqual(self.api.excel_process_preview({'filePath': str(self.root / 'missing.xlsx')})['code'], 0)

    def test_merge_trims_and_deduplicates_after_field_mapping(self):
        first = self.workbook('first.xlsx', [['id', 'name'], ['001', ' Alice ']])
        second = self.workbook('second.xlsx', [['名称', '编号'], ['Alice', '001'], [' Bob ', '002']])
        result = self.api.excel_merge_tables({'tables': [{'path': str(first)}, {'path': str(second),
                                             'fieldMapping': {'id': '编号', 'name': '名称'}}],
                                             'trimText': True, 'deduplicateColumns': ['id', 'name'],
                                             'outputDir': str(self.root / 'out')})
        self.assertEqual(result['code'], 0, result)
        self.assertEqual((result['sourceRows'], result['rows'], result['removedDuplicates']), (3, 2, 1))
        output = load_workbook(result['output'])
        try:
            self.assertEqual(list(output.active.values), [('id', 'name'), ('001', 'Alice'), ('002', 'Bob')])
        finally:
            output.close()

    def test_formula_looking_text_remains_text_after_trimming(self):
        source = self.workbook('text.xlsx', [['_value'], [' =1+2 ']])
        preview = self.api.excel_process_preview({'filePath': str(source), 'trimText': True})
        self.assertEqual(preview['after'][0]['_value'], '=1+2')
        result = self.api.excel_process({'filePath': str(source), 'trimText': True,
                                         'exportCombined': True, 'outputDir': str(self.root / 'out')})
        self.assertEqual(result['code'], 0, result)
        output = load_workbook(result['combinedPath'])
        try:
            self.assertEqual(output.active['A2'].value, '=1+2')
            self.assertEqual(output.active['A2'].data_type, 's')
        finally:
            output.close()


if __name__ == '__main__':
    unittest.main()
